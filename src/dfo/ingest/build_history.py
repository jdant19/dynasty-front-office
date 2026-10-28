"""Build clean history tables in data/processed from the raw files in data/raw.

Run:  PYTHONPATH=src python -m dfo.ingest.build_history

Every output row keeps a `source` column naming the raw file it came from. Parsers never guess an
owner: rows that cannot be mapped get owner = "" and are reported, so they can be fixed in config.
"""
from __future__ import annotations

import csv
import math
import json
import re
import subprocess
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

from dfo.common import PROCESSED, RAW, canon_owner, owners_list, team_owner

LEAGUE = RAW / "league"
XLSX_DRAFTS = LEAGUE / "rookie_drafts_tags_rfa.xlsx"
XLSX_HIST = LEAGUE / "historic_league_info.xlsx"
def colour_outcome(cell) -> str:
    """RFA sheets mark the deciding bid: blue = matched, green = changed teams. Shades vary by year."""
    f = cell.fill
    if f is None or not f.fill_type or f.fgColor is None or f.fgColor.type != "rgb":
        return ""
    rgb = f.fgColor.rgb[-6:]
    r, g, b = (int(rgb[i:i + 2], 16) for i in (0, 2, 4))
    if b > r + 40 and b >= g:
        return "matched"
    if g > r + 30 and g > b + 30:
        return "changed_teams"
    return ""

problems: list[str] = []


def write_csv(name: str, rows: list[dict], fields: list[str]) -> Path:
    PROCESSED.mkdir(parents=True, exist_ok=True)
    path = PROCESSED / f"{name}.csv"
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)
    return path


# --------------------------------------------------------------------------- auctions
AUCTION_LINE = re.compile(
    r"^\s*(\d+)\.\s+(.+?)\s+\((\w*)\s*-\s*([A-Z,/]*)\)\s*\S?\s+\$(\d+)\s+(.+?)\s*$")


def auctions() -> list[dict]:
    rows = []
    for pdf in sorted((RAW / "auctions").glob("yahoo_draft_results_*.pdf")):
        season = int(re.search(r"(\d{4})", pdf.name).group(1))
        text = subprocess.run(["pdftotext", "-layout", str(pdf), "-"], capture_output=True,
                              text=True, check=True).stdout
        seen_picks = set()
        for raw_line in text.splitlines():
            # Yahoo marks kept players with an icon glyph (private-use character) in 2014-2024.
            has_keeper_icon = bool(re.search(r"[-]", raw_line))
            m = AUCTION_LINE.match(re.sub(r"[-]", " ", raw_line))
            if not m:
                continue
            pick, player, nfl, pos, salary, team = m.groups()
            if player.strip() == "--empty--" or int(pick) in seen_picks:
                continue
            seen_picks.add(int(pick))
            if season == 2013:
                kind = "startup_auction"
            elif season == 2025:
                # New Yahoo league: no keeper icons; picks 1-59 are the auction, 60+ are kept players
                # entered team by team.
                kind = "auction" if int(pick) <= 59 else "kept"
            else:
                kind = "kept" if has_keeper_icon else "auction"
            owner = team_owner(team, season) or ""
            if not owner:
                problems.append(f"auction {season}: no owner for team '{team}'")
            rows.append(dict(season=season, pick=int(pick), player=player.strip(), nfl_team=nfl,
                             pos=pos, salary=int(salary), team_name=team.strip(), owner=owner,
                             kind=kind, source=pdf.name))
    return rows


# --------------------------------------------------------------------------- season results
def season_results() -> list[dict]:
    wb = openpyxl.load_workbook(XLSX_HIST, data_only=True)
    ws = wb["Historic League Rankings"]
    data = list(ws.iter_rows(values_only=True))
    hdr = data[0]
    playoff_label = {5: "champion", 4: "runner_up", 3: "semifinal", 0: "missed"}
    out = []
    for r in data[1:]:
        owner = canon_owner(r[0])
        if not owner:
            continue
        for i, h in enumerate(hdr):
            m = re.match(r"(\d{4})\s*-\s*R", str(h or ""))
            if not m or not isinstance(r[i], (int, float)):
                continue
            season = int(m.group(1))
            p = r[i + 1] if i + 1 < len(r) else None
            out.append(dict(season=season, owner=owner, regular_season_finish=11 - int(r[i]),
                            playoff_result=playoff_label.get(int(p)) if isinstance(p, (int, float)) else "",
                            source="historic_league_info.xlsx:Historic League Rankings"))
    return sorted(out, key=lambda x: (x["season"], x["regular_season_finish"]))


# --------------------------------------------------------------------------- rookie drafts
OWNER_DASH = re.compile(r"^\s*([A-Za-z]+)\s*-\s*(.+?)\s*$")
NUM_DASH = re.compile(r"^\s*(\d+)\s*-\s*(.+?)\s*$")


def _draft_entries(ws):
    """Yield (row, col, owner, player) for every pick cell in a rookie draft sheet, any layout."""
    cells = {(c.row, c.column): c.value for row in ws.iter_rows() for c in row if c.value is not None}
    stop_row = min([r for (r, c), v in cells.items() if isinstance(v, str) and
                    re.match(r"(future|picks?\s+20)", v.strip(), re.I)] or [10 ** 6])
    for (r, c), v in sorted(cells.items()):
        if r >= stop_row or not isinstance(v, str):
            continue
        m = OWNER_DASH.match(v)
        if m and canon_owner(m.group(1)):
            yield r, c, canon_owner(m.group(1)), m.group(2)
            continue
        m = NUM_DASH.match(v)
        if m and canon_owner(cells.get((r, c + 1))):
            yield r, c, canon_owner(cells.get((r, c + 1))), m.group(2)
            continue
        if re.match(r"^Pick \d+$", v.strip()):            # 2014 layout: 'Pick N' | owner | player
            o, p = cells.get((r, c + 1)), cells.get((r, c + 2))
            if canon_owner(o) and isinstance(p, str):
                yield r, c, canon_owner(o), p


def rookie_drafts() -> list[dict]:
    rows = []
    seen = set()
    for path in (XLSX_DRAFTS, XLSX_HIST):
        wb = openpyxl.load_workbook(path, data_only=True)
        for ws in wb.worksheets:
            m = re.match(r"(\d{4}) Rookie Draft$", ws.title.strip())
            if not m:
                continue
            season = int(m.group(1))
            if season in seen:
                continue
            entries = list(_draft_entries(ws))
            if not entries:
                continue                                   # e.g. 2027 placeholder sheet
            seen.add(season)
            cols = sorted({c for _, c, _, _ in entries})
            for rnd, col in enumerate(cols[:3], start=1):
                picks = sorted([e for e in entries if e[1] == col], key=lambda e: e[0])
                for i, (_, _, owner, player) in enumerate(picks, start=1):
                    forfeited = "forfeit" in player.lower() or "forefiet" in player.lower()
                    rows.append(dict(season=season, round=rnd, pick=i, overall=(rnd - 1) * 10 + i,
                                     owner=owner, player="" if forfeited else player.strip(),
                                     forfeited=forfeited, salary=({1: 5, 2: 3, 3: 1}[rnd]),
                                     source=f"{path.name}:{ws.title}"))
                if len(picks) != 10:
                    problems.append(f"rookie draft {season} round {rnd}: {len(picks)} picks found")
    return sorted(rows, key=lambda r: (r["season"], r["overall"]))


# --------------------------------------------------------------------------- designations
def designations() -> list[dict]:
    wb = openpyxl.load_workbook(XLSX_DRAFTS, data_only=True)
    rows = []
    for ws in wb.worksheets:
        m = re.match(r"(\d{4})\s*-+\s*RFA and Franchise", ws.title.strip())
        if not m:
            continue
        season = int(m.group(1))
        grid = [[c.value for c in r] for r in ws.iter_rows()]
        header = grid[0]
        kinds = []
        current = None
        for h in header:
            if h and "Restricted" in str(h):
                current = "rfa"
            elif h and "Franchise" in str(h):
                current = "franchise_tag"
            elif h and "Mid" in str(h):
                current = "mid_level_extension"
            kinds.append(current)
        for r in grid[1:]:
            owner = canon_owner(r[0])
            if not owner:
                continue
            for i, v in enumerate(r[1:], start=1):
                if not v or str(v).strip() in ("N/A", "-", "x"):
                    continue
                rows.append(dict(season=season, owner=owner, kind=kinds[i] if i < len(kinds) else "",
                                 player=str(v).strip(), source=f"{XLSX_DRAFTS.name}:{ws.title}"))
    return rows


# --------------------------------------------------------------------------- RFA bids
def rfa_bids() -> list[dict]:
    wb = openpyxl.load_workbook(XLSX_DRAFTS)          # styles needed for outcome colours
    rows = []
    for ws in wb.worksheets:
        m = re.match(r"(\d{4})\s*-?\s*RFA Results", ws.title.strip())
        if not m:
            continue
        season = int(m.group(1))
        all_rows = list(ws.iter_rows())
        hdr_idx = next(i for i, r in enumerate(all_rows)
                       if sum(1 for c in r if canon_owner(c.value)) >= 8)
        hdr = all_rows[hdr_idx]
        owner_cols = {c.column: canon_owner(c.value) for c in hdr if canon_owner(c.value)}
        first_owner_col = min(owner_cols)
        for r in all_rows[hdr_idx + 1:]:
            vals = {c.column: c for c in r if c.value is not None}
            names = [c for col, c in vals.items() if col < first_owner_col and isinstance(c.value, str)]
            nums = [c for col, c in vals.items() if col < first_owner_col and isinstance(c.value, (int, float))]
            if not names or not nums:
                continue
            player = names[0].value.strip()
            start = int(nums[0].value)
            original = next((owner_cols[col] for col, c in vals.items()
                             if col in owner_cols and str(c.value).strip().lower() == "x"), "")
            colored = [(col, c) for col, c in vals.items() if colour_outcome(c)]
            outcome, winner, price = "", "", ""
            if len(colored) > 1:
                problems.append(f"rfa {season} {player}: {len(colored)} coloured cells")
            if colored:
                col, c = colored[0]
                outcome = colour_outcome(c)
                if col in owner_cols:
                    price = int(c.value) if isinstance(c.value, (int, float)) else ""
                    winner = original if outcome == "matched" else owner_cols[col]
                else:                                       # colour on the starting bid: no bids
                    price, winner = start, original
            for col, owner in owner_cols.items():
                c = vals.get(col)
                if c is None or not isinstance(c.value, (int, float)):
                    continue
                rows.append(dict(season=season, player=player, starting_bid=start, original_owner=original,
                                 bidder=owner, bid=int(c.value), outcome=outcome, final_owner=winner,
                                 final_price=price, source=f"{XLSX_DRAFTS.name}:{ws.title}"))
            if not any(c is not None and isinstance(c.value, (int, float)) for col, c in vals.items()
                       if col in owner_cols):
                rows.append(dict(season=season, player=player, starting_bid=start, original_owner=original,
                                 bidder="", bid="", outcome=outcome or "no_bids", final_owner=winner or original,
                                 final_price=price or start, source=f"{XLSX_DRAFTS.name}:{ws.title}"))
            if not outcome:
                problems.append(f"rfa {season} {player}: outcome colour not found")
    return rows


# --------------------------------------------------------------------------- tag benchmarks
def tag_benchmarks() -> list[dict]:
    rows, seen = [], set()
    for path in (XLSX_DRAFTS, XLSX_HIST):
        wb = openpyxl.load_workbook(path, data_only=True)
        for ws in wb.worksheets:
            m = re.match(r"(\d{4}) Rookie Draft$", ws.title.strip())
            if not m:
                continue
            season = int(m.group(1))
            if season > 2026:
                continue                                    # placeholder sheet copies the prior year
            for r in ws.iter_rows(values_only=True):
                vals = [v for v in r if v is not None]
                for i, v in enumerate(vals):
                    if v in ("QB", "RB", "WR", "TE", "K", "D") and i + 1 < len(vals) \
                            and isinstance(vals[i + 1], (int, float)):
                        key = (season, v)
                        if key in seen:
                            continue
                        seen.add(key)
                        nums = []
                        for x in vals[i + 1:i + 7]:
                            if not isinstance(x, (int, float)):
                                break
                            nums.append(x)
                        # Layouts: 6 numbers = average + top 5 (2025-26); 5 = top 5 only (2024);
                        # 1 = average only (2021-23).
                        if len(nums) >= 6:
                            avg, top5 = nums[0], nums[1:6]
                        elif len(nums) == 5:
                            avg, top5 = sum(nums) / 5, nums
                        else:
                            avg, top5 = nums[0], []
                        rows.append(dict(season=season, position="DEF" if v == "D" else v,
                                         sheet_value=round(avg, 2), tag_cost=int(math.floor(avg + 0.5)),
                                         top5_salaries=" ".join(str(int(x)) for x in top5),
                                         source=f"{path.name}:{ws.title}"))
    return sorted(rows, key=lambda r: (r["season"], r["position"]))


# --------------------------------------------------------------------------- trades
def trades() -> list[dict]:
    out = []
    arch = json.load(open(LEAGUE / "old_system_trades_archive.json"))["trades"]
    for i, t in enumerate(arch, start=1):
        for side, other in (("team1", "team2"), ("team2", "team1")):
            for item in t[f"{side}_receives"]:
                out.append(dict(trade_id=f"T{t['season'][:4]}-{i:03d}", season=t["season"],
                                date="", receiver=canon_owner(t[f"{side}_owner"]) or t[f"{side}_owner"],
                                sender=canon_owner(t[f"{other}_owner"]) or t[f"{other}_owner"],
                                asset_type=item["type"],
                                asset=item.get("desc", "") or "", cap_amount=item.get("amount", ""),
                                source="old_system_trades_archive.json"))
    ledger = [json.loads(l) for l in open(LEAGUE / "old_system_ledger" / "events.jsonl")]
    for j, e in enumerate([e for e in ledger if e.get("type") == "trade"], start=1):
        out.append(dict(trade_id=f"T2026-L{j:02d}", season="2026-2027", date=e.get("date", ""),
                        receiver="", sender="", asset_type="summary", asset=e.get("subject", ""),
                        cap_amount="", source="old_system_ledger/events.jsonl (terms in raw file)"))
    inbox = RAW / "inbox" / "trades_2026.json"
    if inbox.exists():
        for k, t in enumerate(json.load(open(inbox)), start=1):
            owners = list(t["sides"])
            for rec in owners:
                sender = [o for o in owners if o != rec][0]
                for a in t["sides"][rec]["receives"]:
                    cap = re.match(r"\$(\d+)", a)
                    out.append(dict(trade_id=f"T2026-J{k:02d}", season=t["season"], date=t["reported"],
                                    receiver=rec, sender=sender, asset_type="cap" if cap else
                                    ("pick" if re.search(r"\b(1st|2nd|3rd)\b", a) else "player"),
                                    asset=a, cap_amount=int(cap.group(1)) if cap else "",
                                    source="inbox/trades_2026.json"))
    return out


# --------------------------------------------------------------------------- 2026 schedule and results
TEAM_2026 = {"Ja'Jah Pachulia": "Jordan", "Who Dey White House": "Ben", "Duck Prescott": "Ryan",
             "REESE MODE": "Matt", "Dayton Triangles": "Qasim", "Hello Fascism": "Johnny",
             "In the Nix of Time": "Zach", "Mike's Swell Team": "Mike", "STEIN DYNASTY": "Russ",
             "Top of the Food Achane": "Max"}
WEEKS_1_9 = {  # from the 2026 matchup screenshots; weeks 10-18 repeat weeks 1-9
    1: [("Jordan", "Ben"), ("Ryan", "Matt"), ("Qasim", "Johnny"), ("Zach", "Mike"), ("Russ", "Max")],
    2: [("Jordan", "Ryan"), ("Qasim", "Matt"), ("Johnny", "Russ"), ("Zach", "Max"), ("Mike", "Ben")],
    3: [("Jordan", "Qasim"), ("Ryan", "Mike"), ("Johnny", "Max"), ("Zach", "Ben"), ("Matt", "Russ")],
    4: [("Jordan", "Russ"), ("Ryan", "Ben"), ("Qasim", "Mike"), ("Johnny", "Zach"), ("Matt", "Max")],
    5: [("Jordan", "Max"), ("Ryan", "Zach"), ("Qasim", "Ben"), ("Johnny", "Matt"), ("Mike", "Russ")],
    6: [("Jordan", "Johnny"), ("Ryan", "Qasim"), ("Zach", "Matt"), ("Mike", "Max"), ("Russ", "Ben")],
    7: [("Jordan", "Matt"), ("Ryan", "Russ"), ("Qasim", "Zach"), ("Johnny", "Mike"), ("Max", "Ben")],
    8: [("Jordan", "Zach"), ("Ryan", "Max"), ("Qasim", "Russ"), ("Johnny", "Ben"), ("Mike", "Matt")],
    9: [("Jordan", "Mike"), ("Ryan", "Johnny"), ("Qasim", "Max"), ("Zach", "Russ"), ("Matt", "Ben")],
}
SCORES = {  # (week, owner) -> points, from screenshots
    1: {"Jordan": 129.16, "Ben": 118.72, "Ryan": 84.60, "Matt": 96.96, "Qasim": 130.96, "Johnny": 150.14,
        "Zach": 145.80, "Mike": 104.26, "Russ": 142.96, "Max": 143.92},
    2: {"Jordan": 117.78, "Ryan": 140.86, "Qasim": 83.32, "Matt": 117.30, "Johnny": 124.76, "Russ": 84.78,
        "Zach": 146.78, "Max": 125.94, "Mike": 78.42, "Ben": 132.58},
    3: {"Jordan": 144.04, "Qasim": 111.58, "Ryan": 101.64, "Mike": 80.16, "Johnny": 104.76, "Max": 80.62,
        "Zach": 161.94, "Ben": 116.64, "Matt": 97.34, "Russ": 132.18},
    4: {"Jordan": 105.66, "Russ": 153.62, "Ryan": 110.50, "Ben": 149.80, "Qasim": 111.48, "Mike": 67.42,
        "Johnny": 174.68, "Zach": 94.08, "Matt": 123.88, "Max": 145.52},
    5: {"Jordan": 160.48, "Max": 135.76, "Ryan": 88.94, "Zach": 173.92, "Qasim": 121.84, "Ben": 68.76,
        "Johnny": 138.54, "Matt": 86.24, "Mike": 115.00, "Russ": 170.60},
    6: {"Jordan": 125.74, "Johnny": 113.72},
    7: {"Jordan": 128.36, "Matt": 123.02, "Ryan": 117.18, "Russ": 162.78, "Qasim": 99.42, "Zach": 129.70,
        "Johnny": 141.82, "Mike": 97.14, "Max": 109.38, "Ben": 82.90},
}
STANDINGS_AFTER_W7 = {  # owner: (wins, losses, points_for, points_against), screenshot 2026-10-27
    "Johnny": (6, 1, 948.42, 699.56), "Russ": (5, 2, 981.04, 788.04), "Zach": (5, 2, 953.70, 825.54),
    "Jordan": (5, 2, 911.22, 897.28), "Max": (4, 3, 861.44, 833.86), "Qasim": (3, 4, 774.90, 759.98),
    "Matt": (3, 4, 760.40, 814.00), "Ben": (2, 5, 753.58, 845.36), "Ryan": (2, 5, 726.34, 897.70),
    "Mike": (0, 7, 614.50, 924.22)}


def results_2026() -> list[dict]:
    scores = {w: dict(s) for w, s in SCORES.items()}
    # Week 6 scores for the other eight teams: points_for after week 7 minus weeks 1-5 and 7.
    for owner, (_, _, pf, _) in STANDINGS_AFTER_W7.items():
        if owner not in scores[6]:
            scores[6][owner] = round(pf - sum(scores[w][owner] for w in (1, 2, 3, 4, 5, 7)), 2)
    rows = []
    for week in range(1, 19):
        base = week if week <= 9 else week - 9
        for a, b in WEEKS_1_9[base]:
            sa, sb = scores.get(week, {}).get(a), scores.get(week, {}).get(b)
            derived = week == 6 and (a not in SCORES[6] or b not in SCORES[6])
            rows.append(dict(season=2026, week=week, team_a=a, team_b=b,
                             score_a=sa if sa is not None else "", score_b=sb if sb is not None else "",
                             winner=(a if sa > sb else b) if sa is not None and sb is not None else "",
                             basis=("derived_from_standings" if derived else "screenshot")
                             if sa is not None else "scheduled",
                             source="screenshots/2026" + ("" if week <= 9 else " (repeat of week %d)" % base)))
    return rows


# --------------------------------------------------------------------------- main
def main() -> None:
    outputs = {
        "auctions": (auctions(), ["season", "pick", "player", "nfl_team", "pos", "salary", "team_name",
                                  "owner", "kind", "source"]),
        "season_results": (season_results(), ["season", "owner", "regular_season_finish",
                                              "playoff_result", "source"]),
        "rookie_drafts": (rookie_drafts(), ["season", "round", "pick", "overall", "owner", "player",
                                            "forfeited", "salary", "source"]),
        "designations": (designations(), ["season", "owner", "kind", "player", "source"]),
        "rfa_bids": (rfa_bids(), ["season", "player", "starting_bid", "original_owner", "bidder", "bid",
                                  "outcome", "final_owner", "final_price", "source"]),
        "tag_benchmarks": (tag_benchmarks(), ["season", "position", "sheet_value", "tag_cost",
                                              "top5_salaries", "source"]),
        "trades": (trades(), ["trade_id", "season", "date", "receiver", "sender", "asset_type", "asset",
                              "cap_amount", "source"]),
        "results_2026": (results_2026(), ["season", "week", "team_a", "team_b", "score_a", "score_b",
                                          "winner", "basis", "source"]),
    }
    for name, (rows, fields) in outputs.items():
        p = write_csv(name, rows, fields)
        print(f"{name:16s} {len(rows):5d} rows -> {p.relative_to(PROCESSED.parents[1])}")
    if problems:
        print(f"\n{len(problems)} problem(s):")
        for msg in Counter(problems):
            print("  -", msg)


if __name__ == "__main__":
    main()
